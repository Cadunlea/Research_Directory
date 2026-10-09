"""Results summary for Dr. Sazonov, week of 6 October 2026.

Every number below is copied from the training logs of that week; nothing is
estimated. Arial 11, grayscale tables, US Letter.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT, INK = "Arial", RGBColor(0x22, 0x22, 0x22)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(0.9))

def style_font(style, size, bold=None):
    style.font.name, style.font.size, style.font.color.rgb = FONT, Pt(size), INK
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(key), FONT)
    # Theme fonts override explicit ones; headings would stay in the theme face.
    for key in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(key)) is not None:
            del fonts.attrib[qn(key)]
    # Grayscale only: drop the theme's coloured rule under the title.
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)

style_font(doc.styles["Normal"], 11)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11)):
    style_font(doc.styles[name], size, True)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(4)

def para(text, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic, r.font.size = italic, Pt(size)
    return p

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def table(caption, header, rows, widths, note=None, bold_row=None):
    cap = doc.add_paragraph()
    run = cap.add_run(caption)
    run.bold = True
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "A6A6A6")
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for r, values in enumerate([header] + rows):
        for c, value in enumerate(values):
            cell = t.rows[r].cells[c]
            cell.width = Inches(widths[c])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if c > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(11)
            run.bold = r == 0 or (bold_row is not None and r == bold_row + 1)
            if r == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EFEFEF")
                cell._tc.get_or_add_tcPr().append(shd)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    if note:
        n = para(note, italic=True, size=11)
        n.paragraph_format.space_before = Pt(3)

# --------------------------------------------------------------------------
doc.add_paragraph("Food-Intake Detection from AIM-2 Signals: Experiment Results", style="Title")
para("Caelan Dunlea  ·  9 October 2026")

doc.add_heading("Summary", level=1)
bullets([
    "Under leave-one-subject-out validation (43 participants), the model reaches pooled F1 0.856 and "
    "accuracy 0.887; participant-level 5-fold cross-validation gives F1 0.864 and accuracy 0.893.",
    "The input representation has the largest effect measured: replacing the time series with an FFT "
    "magnitude spectrum lowers F1 from 0.858 to 0.796 and raises false alarms from 32 to 81 per hour "
    "of non-eating.",
    "8 s windows and the full 128 Hz sampling rate perform best. 32 Hz costs about 0.02 F1, almost "
    "entirely in recall.",
    "Training settings, context windows and a Transformer layer change F1 by 0.01 or less.",
])

doc.add_heading("Dataset and protocol", level=1)
table("Table 1. Dataset", ["Item", "Value"], [
    ["Annotated sessions", "61 (breakfast 8, lunch 18, dinner 12, snack 23)"],
    ["Participants", "43 (18 with two sessions)"],
    ["Annotated recording", "24.2 h, of which 9.6 h eating"],
    ["Evaluation windows", "10,891 non-overlapping 8 s windows"],
    ["Sensor input", "3-axis accelerometer and optical sensor, 128 Hz"],
], [2.0, 4.7])
para("All folds hold out whole participants with all of their sessions. Each comparison below uses "
     "the same folds; only the stated setting changes. Unless noted, results are pooled over all "
     "held-out windows from participant-level 5-fold cross-validation, and the decision threshold "
     "is chosen on validation participants, never on the test participants.")

doc.add_heading("Main result", level=1)
table("Table 2. Current model: leave-one-subject-out vs. 5-fold cross-validation",
      ["Metric", "LOSO (43 folds)", "5-fold"], [
    ["F1", "0.856", "0.864"],
    ["Accuracy", "0.887", "0.893"],
    ["Balanced accuracy", "0.881", "0.887"],
    ["Precision", "0.860", "0.867"],
    ["Recall", "0.853", "0.861"],
    ["Specificity", "0.909", "0.914"],
    ["False alarms per hour of non-eating", "41.1", "38.8"],
    ["Per-participant F1, mean ± SD", "0.836 ± 0.117", "0.840 ± 0.129"],
    ["Per-participant F1, median", "0.867", "0.881"],
], [3.0, 1.85, 1.85],
note="Lowest participants under LOSO: AIM147560 (F1 0.452, 46 test windows), AIM112400 (0.525), "
     "AIM132097 (0.545).")

doc.add_heading("Input representation", level=1)
table("Table 3. Time series vs. FFT magnitude spectrum (same network)",
      ["Metric", "Time series", "FFT spectrum"], [
    ["F1", "0.858", "0.796"],
    ["Accuracy", "0.891", "0.829"],
    ["Precision", "0.884", "0.754"],
    ["Recall", "0.834", "0.843"],
    ["False alarms per hour of non-eating", "32.4", "81.0"],
    ["Per-participant F1, mean", "0.835", "0.756"],
], [3.0, 1.85, 1.85],
note="Earlier training settings (learning rate 0.001, early stopping on validation loss). The time "
     "series scored higher in all five folds.")

doc.add_heading("Sampling rate", level=1)
table("Table 4. Sampling rate (anti-aliased downsampling from 128 Hz)",
      ["Rate", "F1", "Accuracy", "Precision", "Recall", "False alarms/h"], [
    ["128 Hz", "0.864", "0.893", "0.867", "0.861", "38.8"],
    ["64 Hz", "0.851", "0.884", "0.865", "0.838", "38.6"],
    ["32 Hz", "0.841", "0.879", "0.871", "0.814", "35.6"],
    ["16 Hz", "0.842", "0.878", "0.865", "0.821", "37.8"],
], [1.1, 1.0, 1.1, 1.1, 1.1, 1.3], bold_row=0,
note="Lower rates reduce recall; precision and false alarms are unchanged. 16 Hz matches 32 Hz.")

doc.add_heading("Window length", level=1)
table("Table 5. Window length",
      ["Window", "Windows", "F1", "Precision", "Recall", "Per-participant F1"], [
    ["2 s", "43,708", "0.829", "0.834", "0.825", "0.804"],
    ["4 s", "21,828", "0.845", "0.867", "0.824", "0.821"],
    ["8 s", "10,891", "0.864", "0.867", "0.861", "0.840"],
    ["16 s", "5,422", "0.822", "0.881", "0.770", "0.771"],
], [0.9, 1.1, 0.9, 1.1, 1.0, 1.7], bold_row=2,
note="False alarms per hour are not compared across window lengths, since they count windows of "
     "different durations. 16 s windows lose recall because they often mix eating with pauses.")

doc.add_heading("Architecture and training options", level=1)
table("Table 6. Additions to the current model",
      ["Variant", "F1", "Precision", "Recall", "False alarms/h", "Time per fold"], [
    ["Current model", "0.864", "0.867", "0.861", "38.8", "~55 s"],
    ["+ 8 s context each side (with marker)", "0.853", "0.850", "0.857", "44.5", "~180 s"],
    ["+ 1 Transformer layer", "0.862", "0.876", "0.848", "35.3", "~105 s"],
], [2.6, 0.75, 0.9, 0.75, 1.0, 0.9], bold_row=0)

table("Table 7. Training settings",
      ["Setting", "F1", "Accuracy", "Recall", "False alarms/h"], [
    ["Learning rate 0.001, stop on validation loss", "0.858", "0.891", "0.834", "32.4"],
    ["Learning rate 0.001, stop on validation AUC", "0.858", "0.891", "0.835", "32.5"],
    ["Learning rate 0.0003, stop on validation AUC", "0.864", "0.893", "0.861", "38.8"],
    ["Augmentation, stop on validation AUC", "0.862", "0.893", "0.846", "34.3"],
    ["Augmentation, learning rate 0.0003", "0.854", "0.886", "0.842", "38.5"],
], [3.1, 0.85, 0.95, 0.85, 1.0], bold_row=2,
note="At learning rate 0.001, early stopping kept the first-epoch model in three of five folds; "
     "0.0003 moves the best epoch to 2–6. Differences of 0.01 F1 or less are not yet confirmed as "
     "significant; paired per-participant tests and a three-seed repeat are in progress.")

doc.add_heading("In progress", level=1)
bullets([
    "Weight of the chew-count head (0, 0.1, 0.2, 0.5, 1.0).",
    "Decision threshold set for a false-alarm budget (15 per hour) and for F0.5.",
    "Ablations of channel attention and attention pooling on all 61 sessions.",
    "Published FCN and ResNet architectures (as compared by Ghosh and Sazonov, 2022) on the same folds.",
    "Three-seed repeat and paired per-participant statistics for every comparison.",
])

out = "/home/user/Research_Directory/docs/progress_reports/results_summary_2026-10-09.docx"
doc.save(out)
print("saved", out)
