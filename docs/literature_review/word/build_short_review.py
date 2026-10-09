"""Shortened literature review and model description, as a Word document.

Cites only the ten papers on the literature-review slides. Arial 11,
grayscale tables in the style of the weekly results summary, US Letter.
Run from this folder; needs ../figure_architecture.png.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIGURE = HERE.parent / "figure_architecture.png"
OUT = HERE.parent / "literature_review_short.docx"
FONT, INK = "Arial", RGBColor(0x22, 0x22, 0x22)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(0.9))


def style_font(style, size, bold=None):
    style.font.name, style.font.size, style.font.color.rgb = FONT, Pt(size), INK
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(key), FONT)
    for key in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(key)) is not None:
            del fonts.attrib[qn(key)]
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)


style_font(doc.styles["Normal"], 11)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11)):
    style_font(doc.styles[name], size, True)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(4)
style_font(doc.styles["List Bullet"], 11)


def para(text, italic=False):
    p = doc.add_paragraph()
    p.add_run(text).italic = italic
    return p


def caption(text):
    p = doc.add_paragraph()
    p.add_run(text).bold = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    return p


def table(header, rows, widths, centre_from=1):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "A6A6A6")
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for r, values in enumerate([header] + rows):
        for c, value in enumerate(values):
            cell = t.rows[r].cells[c]
            cell.width = Inches(widths[c])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if centre_from is not None and c >= centre_from:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(11)
            run.bold = r == 0
            if r == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EFEFEF")
                cell._tc.get_or_add_tcPr().append(shd)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# ---------------------------------------------------------------------------
doc.add_paragraph("Food-Intake Detection from AIM-2 Sensor Signals: "
                  "Literature Review and Model Description", style="Title")
para("Caelan Dunlea  ·  University of Alabama  ·  9 October 2026")

doc.add_heading("1  Literature Review", level=1)
para("Self-reported food intake is inaccurate and cannot capture how people eat within a meal, "
     "such as eating rate or pauses (Diou et al., 2022). Wearable sensors that detect chewing or "
     "bites address this. The Automatic Ingestion Monitor version 2 (AIM-2) places a "
     "temporalis-muscle sensor and a 3-axis accelerometer on eyeglasses; its original detector "
     "extracted 38 hand-crafted features per signal and classified 10 s epochs with a support "
     "vector machine, reaching F1 0.818 under leave-one-subject-out validation on 30 participants "
     "(Doulah et al., 2021). Feature-engineered detectors of this kind are limited by the features "
     "chosen by hand and are sensitive to variation between people. Table 1 summarises the studies "
     "reviewed.")

caption("Table 1. Reviewed studies")
table(["Reference", "Technique", "Findings"], [
    ["Doulah et al., 2021", "AIM-2; hand-crafted features and SVM, 10 s windows",
     "F1 0.818 (leave-one-out, 30 participants)"],
    ["Ghosh and Sazonov, 2022", "Five deep networks on AIM-2 optical and accelerometer signals",
     "ResNet best, balanced accuracy 93.5%"],
    ["Ghosh et al., 2024", "Wavelet image of accelerometer, fused with food images",
     "Fused F1 0.808; fusion reduces false positives"],
    ["Kyritsis et al., 2021", "CNN + LSTM trained end to end on wrist IMU",
     "Bite F1 0.923 (leave-one-out)"],
    ["Diou et al., 2022", "Review: CNN on in-ear audio for chewing, smartwatch for bites",
     "CNN F1 0.908 vs. 0.748 for an SVM"],
    ["Stankoski et al., 2021", "Hard negative examples selected for training",
     "F1 0.82 on unseen people; fewer false positives"],
    ["Wang et al., 2024", "Temporal CNN + multi-head self-attention, wrist IMU",
     "61 participants; eating-speed error 11%"],
    ["Vedovelli et al., 2026", "Transformer on smartwatch sensors",
     "Sensitivity 0.69, specificity 0.50 (leave-one-out)"],
    ["Hannun et al., 2019", "34-layer CNN on raw ECG",
     "F1 0.837, above cardiologists (0.780)"],
    ["Ikram et al., 2025", "Transformer on ECG",
     "Accuracy 97.1%; split by beats, not patients"],
], [1.65, 2.75, 2.3], centre_from=None)

doc.add_heading("1.1  Learned features", level=2)
para("Networks trained on the signal itself have outperformed hand-crafted features. On wrist "
     "inertial data, an end-to-end CNN and LSTM reached a bite-detection F1 of 0.923 and "
     "outperformed hand-engineered methods (Kyritsis et al., 2021). For chewing, a CNN on in-ear "
     "audio reached F1 0.908 against 0.748 for an SVM, although such models “require "
     "significantly larger volumes of labeled data” (Diou et al., 2022, p. 8). On AIM-2 itself, "
     "five deep architectures were compared on the optical and accelerometer time series, with a "
     "residual network best at 93.5% balanced accuracy (Ghosh and Sazonov, 2022). Outside eating, "
     "a CNN applied to raw ECG exceeded the average cardiologist (Hannun et al., 2019). In this "
     "literature, a raw input means one without hand-engineered features; light preprocessing such "
     "as normalisation is still applied. Ghosh et al. (2024) instead converted the accelerometer "
     "signal to a wavelet image. None of the reviewed studies compares time-series and spectral "
     "input on the same network; that comparison is made in this work.")

doc.add_heading("1.2  Attention and Transformers", level=2)
para("Attention lets a network weight the parts of its input that carry evidence. A temporal CNN "
     "with multi-head self-attention measured eating speed on 61 participants (Wang et al., 2024). "
     "A Transformer encoder on smartwatch data, however, reached a sensitivity of 0.69 but a "
     "specificity of only 0.50 under leave-one-subject-out validation (Vedovelli et al., 2026), and "
     "the 97.1% accuracy reported for a Transformer on ECG comes from a split of beats rather than "
     "patients (Ikram et al., 2025). This work therefore uses light attention on a CNN by default "
     "and evaluates a Transformer layer as a separate option.")

doc.add_heading("1.3  False positives and evaluation", level=2)
para("False positives are the main weakness of eating detectors. Walking has a rhythm similar to "
     "chewing, 1 to 3 Hz, and causes chewing classifiers “to yield high false-positive "
     "detections”, which accelerometer data can be used to remove (Diou et al., 2022, p. 8). "
     "Selecting hard negative examples for training (Stankoski et al., 2021) and fusing sensor and "
     "image classifiers (Ghosh et al., 2024) both reduced false positives, and a detector can be "
     "tuned “towards higher recall or precision, or a balanced mode, based on the needs of each "
     "use-case” (Diou et al., 2022, p. 8). The intake studies reviewed evaluate on held-out "
     "people (Doulah et al., 2021; Kyritsis et al., 2021; Vedovelli et al., 2026), which is the "
     "protocol adopted here.")

doc.add_heading("1.4  Contribution of this work", level=2)
for item in [
    "A compact convolutional model with channel attention, temporal attention pooling and an "
    "auxiliary chew-count head, trained on normalised AIM-2 optical and accelerometer signals from "
    "61 annotated sessions (43 participants) and evaluated leave-one-subject-out.",
    "A direct comparison of time-series and spectral input on the same network: the time series "
    "reaches F1 0.858 against 0.796 and reduces false alarms from 81 to 32 per hour of non-eating.",
    "Controlled experiments that change one design choice at a time on identical participant-level "
    "folds, including negative results reported with their mechanism.",
]:
    doc.add_paragraph(item, style="List Bullet")

# ---------------------------------------------------------------------------
doc.add_heading("2  Model Description", level=1)
para("The model is a one-dimensional convolutional network that reads 8 s windows of the "
     "normalised AIM-2 signals and outputs the probability that the window contains eating "
     "(Figure 1). Three convolution blocks learn their own filters from the time series; a "
     "squeeze-and-excitation gate reweights the feature channels in each window; an attention "
     "pooling layer weights the time steps that carry the evidence; and a chew-count head, used "
     "only during training, adds supervision from the chew annotations. The network has 47,282 "
     "parameters at inference.")

pic = doc.add_paragraph()
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic.add_run().add_picture(str(FIGURE), width=Inches(4.6))
para("Figure 1. Model architecture. Shaded blocks are the two attention mechanisms; the dashed "
     "chew-count head is used only during training. Output shapes are time steps × features.",
     italic=True)

caption("Table 2. Layer dimensions")
table(["Layer", "Configuration", "Output", "Parameters"], [
    ["Input", "Acc X, Y, Z, optical", "1024 × 4", "–"],
    ["Conv block 1", "32 filters, kernel 9, stride 2, max-pool 2", "256 × 32", "1,312"],
    ["Conv block 2", "64 filters, kernel 9, stride 2, max-pool 2", "64 × 64", "18,752"],
    ["Conv block 3", "64 filters, kernel 5", "64 × 64", "20,800"],
    ["Channel attention", "64 → 16 → 64, sigmoid", "64 × 64", "2,128"],
    ["Attention pooling", "Score per step, softmax over time", "64", "65"],
    ["Dense", "64 units, ReLU, dropout 0.3", "64", "4,160"],
    ["Eating head", "1 unit, sigmoid", "1", "65"],
    ["Total at inference", "", "", "47,282"],
    ["Chew-count head", "32 → 1, softplus (training only)", "1", "2,113"],
], [1.55, 2.95, 1.0, 1.2], centre_from=2)

caption("Table 3. Design decisions")
table(["Component", "Design", "Rationale"], [
    ["Window", "8 s at 128 Hz, accelerometer X, Y, Z and optical",
     "Chewing lies at 1–3 Hz, so 8 s holds at least eight chews (Diou et al., 2022); best of "
     "2–16 s tested here"],
    ["Labels", "Eating = bout or bite; window is eating if over half its annotated samples are; "
     "unrecorded time dropped",
     "Treating unannotated time as non-eating would bury the real labels"],
    ["Alignment", "No time shift",
     "Measured lag of 0.1 s between sensor and annotations (this work)"],
    ["Input", "Normalised time series; no spectrum, no hand-crafted features",
     "Learned features outperformed hand-crafted ones (Ghosh and Sazonov, 2022; Kyritsis et al., "
     "2021); time series vs. FFT tested here"],
    ["Normalisation", "Per window and channel: mean removed, z-scored; gaps set to 0",
     "Removes each participant’s sensor baseline"],
    ["Convolution", "Three 1D conv blocks with batch norm and ReLU",
     "CNNs on the raw signal (Ghosh and Sazonov, 2022; Hannun et al., 2019); each output step "
     "covers about one chew cycle"],
    ["Channel attention", "Squeeze-and-excitation gate",
     "Walking resembles chewing; the accelerometer separates the two (Diou et al., 2022)"],
    ["Temporal pooling", "Attention pooling over 64 time steps",
     "Averaging would dilute a 1–2 s bite across 8 s"],
    ["Chew-count head", "Regression of chews per window, loss weight 0.2, training only",
     "Deep models need large amounts of labelled data (Diou et al., 2022); design specific to this "
     "work"],
    ["Training", "50% overlapping windows; Adam, learning rate 0.0003; early stopping on "
     "validation AUC; class-balanced weights",
     "At 0.001 the validation optimum fell at the first epoch; eating share ranges 0–86% by "
     "session"],
    ["Validation", "Leave-one-subject-out, 43 folds",
     "Standard in intake detection (Doulah et al., 2021; Kyritsis et al., 2021); sample-level "
     "splits overstate accuracy (Ikram et al., 2025)"],
    ["Threshold", "Chosen on validation participants; best F1 or a false-alarm budget",
     "Detectors can be tuned towards recall or precision by use case (Diou et al., 2022)"],
    ["Post-processing", "None",
     "Median smoothing lowered F1: meals are punctuated by pauses"],
], [1.35, 2.55, 2.8], centre_from=None)

doc.add_heading("2.1  Dataset and Performance", level=2)
para("The dataset is 61 annotated eating sessions (breakfast, lunch, dinner and snack) from 43 "
     "participants, 2 to 55 minutes each: 24.2 hours of annotated recording, of which 9.6 hours is "
     "eating. Each leave-one-subject-out fold holds out one participant with all of their sessions.")

caption("Table 4. Performance (pooled over 10,891 held-out windows)")
table(["Metric", "Leave-one-subject-out", "5-fold"], [
    ["F1", "0.856", "0.864"],
    ["Accuracy", "0.887", "0.893"],
    ["Precision", "0.860", "0.867"],
    ["Recall", "0.853", "0.861"],
    ["False alarms per hour of non-eating", "41.1", "38.8"],
    ["Per-participant F1, median", "0.867", "0.881"],
], [3.1, 1.9, 1.7])

para("The estimate is not sensitive to the validation scheme. Training settings change F1 by no more "
     "than 0.006, while replacing the time series with a magnitude spectrum lowers F1 from 0.858 to "
     "0.796 and raises false alarms from 32 to 81 per hour (5-fold). The variation between "
     "participants is larger than any difference between training configurations.")

# ---------------------------------------------------------------------------
doc.add_heading("References", level=1)
REFS = [
    "Diou, C., Kyritsis, K., Papapanagiotou, V., & Sarafis, I. (2022). Intake monitoring in "
    "free-living conditions: Overview and lessons we have learned. Appetite, 176, 106096. "
    "https://doi.org/10.1016/j.appet.2022.106096",
    "Doulah, A., Ghosh, T., Hossain, D., Imtiaz, M. H., & Sazonov, E. (2021). “Automatic "
    "Ingestion Monitor Version 2” – A novel wearable device for automatic food intake "
    "detection and passive capture of food images. IEEE Journal of Biomedical and Health "
    "Informatics, 25(2), 568–576. https://doi.org/10.1109/JBHI.2020.2995473",
    "Ghosh, T., & Sazonov, E. (2022). A comparative study of deep learning algorithms for detecting "
    "food intake. 44th Annual International Conference of the IEEE Engineering in Medicine and "
    "Biology Society (EMBC), 2993–2996. https://doi.org/10.1109/EMBC48229.2022.9871278",
    "Ghosh, T., Han, Y., Raju, V., Hossain, D., McCrory, M. A., Higgins, J., Boushey, C., Delp, E. "
    "J., & Sazonov, E. (2024). Integrated image and sensor-based food intake detection in "
    "free-living. Scientific Reports, 14, 1665. https://doi.org/10.1038/s41598-024-51687-3",
    "Hannun, A. Y., Rajpurkar, P., Haghpanahi, M., Tison, G. H., Bourn, C., Turakhia, M. P., & Ng, "
    "A. Y. (2019). Cardiologist-level arrhythmia detection and classification in ambulatory "
    "electrocardiograms using a deep neural network. Nature Medicine, 25(1), 65–69. "
    "https://doi.org/10.1038/s41591-018-0268-3",
    "Ikram, S., Ikram, A., Singh, H., Ali Awan, M. D., Naveed, S., De la Torre Díez, I., "
    "Gongora, H. F., & Candelaria Chio Montero, T. (2025). Transformer-based ECG classification for "
    "early detection of cardiac arrhythmias. Frontiers in Medicine, 12, 1600855. "
    "https://doi.org/10.3389/fmed.2025.1600855",
    "Kyritsis, K., Diou, C., & Delopoulos, A. (2021). A data driven end-to-end approach for "
    "in-the-wild monitoring of eating behavior using smartwatches. IEEE Journal of Biomedical and "
    "Health Informatics, 25(1), 22–34.",
    "Stankoski, S., Jordan, M., Gjoreski, H., & Luštrek, M. (2021). Smartwatch-based eating "
    "detection: Data selection for machine learning from imbalanced data with imperfect labels. "
    "Sensors, 21(5), 1902. https://doi.org/10.3390/s21051902",
    "Vedovelli, L., Bhuyan, M. J., Lanera, C., Baldi, I., & Gregori, D. (2026). Recognition of "
    "eating episodes via commercial smartwatch sensors. PLOS Digital Health. "
    "https://doi.org/10.1371/journal.pdig.0001539",
    "Wang, C., Kumar, T. S., De Raedt, W., Camps, G., Hallez, H., & Vanrumste, B. (2024). Eating "
    "speed measurement using wrist-worn IMU sensors towards free-living environments. IEEE Journal "
    "of Biomedical and Health Informatics, 28(10), 5816–5828.",
]
for ref in REFS:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)

doc.save(OUT)
print("saved", OUT)
